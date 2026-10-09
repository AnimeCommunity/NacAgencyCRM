from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import override_settings
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from clientes.models import Cliente
from cotizaciones.models import Cotizacion
from interacciones.models import Interaccion
from marketing.models import ConfigSMTP, PlantillaMensaje
from proyectos.models import Proyecto


User = get_user_model()


class MarketingAndExportTests(APITestCase):
    def setUp(self):
        self.sales = User.objects.create_user(
            username="sales", password="StrongPass!2026", role="sales"
        )
        self.production = User.objects.create_user(
            username="production", password="StrongPass!2026", role="production"
        )
        Cliente.objects.create(
            nombre="ACME", email="acme@example.com", origen="web"
        )

    def test_sales_export_returns_valid_xlsx(self):
        self.client.force_authenticate(self.sales)
        response = self.client.get("/api/marketing/export-excel/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook["Reporte de Clientes"]["A2"].value, "ACME")

    def test_production_cannot_manage_smtp_configuration(self):
        self.client.force_authenticate(self.production)
        response = self.client.post(
            "/api/marketing-config/",
            {
                "email_usuario": "smtp@example.com",
                "servidor_host": "smtp.example.com",
                "puerto": 587,
                "use_tls": True,
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_duplicate_smtp_configuration_is_rejected(self):
        admin = User.objects.create_user(
            username="admin-smtp-duplicate", password="StrongPass!2026", role="admin"
        )
        ConfigSMTP.objects.create(
            email_usuario="first@example.com", servidor_host="smtp.example.com"
        )
        self.client.force_authenticate(admin)

        response = self.client.post(
            "/api/marketing-config/",
            {
                "email_usuario": "second@example.com",
                "servidor_host": "smtp.example.com",
                "puerto": 587,
                "use_tls": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_smtp_singleton_is_enforced_by_database(self):
        ConfigSMTP.objects.create(
            email_usuario="first-db@example.com", servidor_host="smtp.example.com"
        )

        with self.assertRaises(IntegrityError), transaction.atomic():
            ConfigSMTP.objects.create(
                email_usuario="second-db@example.com", servidor_host="smtp.example.com"
            )

    @override_settings(
        CRM_SMTP_ENCRYPTION_KEY="MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=",
    )
    def test_admin_can_store_encrypted_smtp_password_from_api(self):
        admin = User.objects.create_user(
            username="admin-smtp", password="StrongPass!2026", role="admin"
        )
        self.client.force_authenticate(admin)

        response = self.client.post(
            "/api/marketing-config/",
            {
                "email_usuario": "smtp@example.com",
                "email_password": "must-not-be-stored",
                "servidor_host": "smtp.example.com",
                "puerto": 587,
                "use_tls": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        config = ConfigSMTP.objects.get()
        self.assertNotEqual(config.email_password, "")
        self.assertNotEqual(config.email_password, "must-not-be-stored")
        self.assertNotIn("email_password", response.data)
        self.assertTrue(response.data["password_configurada"])

    @override_settings(
        CRM_SMTP_ENCRYPTION_KEY="MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=",
    )
    def test_smtp_update_without_password_preserves_existing_secret(self):
        admin = User.objects.create_user(
            username="admin-smtp-preserve", password="StrongPass!2026", role="admin"
        )
        self.client.force_authenticate(admin)
        created = self.client.post(
            "/api/marketing-config/",
            {
                "email_usuario": "smtp@example.com",
                "email_password": "secret-value",
                "servidor_host": "smtp.example.com",
                "puerto": 587,
                "use_tls": True,
            },
            format="json",
        )
        original_ciphertext = ConfigSMTP.objects.get().email_password

        response = self.client.patch(
            f"/api/marketing-config/{created.data['id']}/",
            {"puerto": 465},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        config = ConfigSMTP.objects.get()
        self.assertEqual(config.email_password, original_ciphertext)
        self.assertTrue(response.data["password_configurada"])

    @override_settings(CRM_SMTP_ENCRYPTION_KEY="")
    def test_smtp_configuration_reports_missing_encryption_key_safely(self):
        admin = User.objects.create_user(
            username="admin-smtp-no-key", password="StrongPass!2026", role="admin"
        )
        self.client.force_authenticate(admin)

        response = self.client.post(
            "/api/marketing-config/",
            {
                "email_usuario": "smtp@example.com",
                "email_password": "secret-value",
                "servidor_host": "smtp.example.com",
                "puerto": 587,
                "use_tls": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("email_password", response.data)

    @override_settings(
        CRM_SMTP_ENCRYPTION_KEY="MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=",
        CRM_SMTP_ALLOWED_HOSTS=["smtp.example.com"],
    )
    @patch("marketing.views.get_connection")
    def test_admin_can_test_smtp_connection_without_exposing_password(self, get_connection):
        admin = User.objects.create_user(
            username="admin-smtp-test", password="StrongPass!2026", role="admin"
        )
        self.client.force_authenticate(admin)
        created = self.client.post(
            "/api/marketing-config/",
            {
                "email_usuario": "smtp@example.com",
                "email_password": "secret-value",
                "servidor_host": "smtp.example.com",
                "puerto": 587,
                "use_tls": True,
            },
            format="json",
        )
        connection = MagicMock()
        get_connection.return_value = connection

        response = self.client.post(
            f"/api/marketing-config/{created.data['id']}/test-connection/"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["status"], "Conexión SMTP verificada.")
        get_connection.assert_called_once_with(
            backend="django.core.mail.backends.smtp.EmailBackend",
            host="smtp.example.com",
            port=587,
            username="smtp@example.com",
            password="secret-value",
            use_tls=True,
            timeout=10,
        )
        connection.open.assert_called_once_with()
        connection.close.assert_called_once_with()

    @override_settings(
        CRM_SMTP_PASSWORD="test-password",
        CRM_SMTP_ALLOWED_HOSTS=["smtp.example.com"],
    )
    def test_campaign_rejects_smtp_host_outside_allowlist(self):
        self.client.force_authenticate(self.sales)
        cliente = Cliente.objects.get(email="acme@example.com")
        plantilla = PlantillaMensaje.objects.create(
            nombre="Seguimiento",
            tipo="email",
            asunto="Hola",
            contenido="Hola {nombre}",
        )
        ConfigSMTP.objects.create(
            email_usuario="smtp@example.com",
            email_password="legacy-password",
            servidor_host="127.0.0.1",
            puerto=25,
            use_tls=False,
        )

        response = self.client.post(
            "/api/marketing/enviar-campana/",
            {"cliente_id": cliente.id, "plantilla_id": plantilla.id},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("permitido", response.data["error"])


class ManagerialReportTests(APITestCase):
    def setUp(self):
        self.sales = User.objects.create_user(
            username="sales-report", password="StrongPass!2026", role="sales"
        )
        self.customer = Cliente.objects.create(
            nombre="Cliente reportes", email="reports@example.com"
        )
        self.project = Proyecto.objects.create(
            cliente=self.customer,
            nombre="Evento",
            descripcion="Evento",
            tipo_evento="show",
            fecha_inicio=date.today(),
            fecha_fin=date.today() + timedelta(days=1),
            presupuesto_estimado=Decimal("500.00"),
            responsable=self.sales,
        )
        Cotizacion.objects.create(
            projecto=self.project,
            numero="A-1",
            subtotal=Decimal("100.00"),
            impuestos=Decimal("19.00"),
            total=Decimal("119.00"),
            estado="aceptada",
            fecha_vencimiento=date.today() + timedelta(days=10),
        )
        Cotizacion.objects.create(
            projecto=self.project,
            numero="R-1",
            subtotal=Decimal("900.00"),
            impuestos=Decimal("171.00"),
            total=Decimal("1071.00"),
            estado="rechazada",
            fecha_vencimiento=date.today() + timedelta(days=10),
        )

    def test_production_can_read_managerial_report(self):
        production = User.objects.create_user(
            username="production-report", password="StrongPass!2026", role="production"
        )
        self.client.force_authenticate(production)

        response = self.client.get("/api/reports/gerencial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_report_counts_projects_and_accepted_revenue_once(self):
        self.client.force_authenticate(self.sales)
        response = self.client.get("/api/reports/gerencial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        metric = response.data["ingresos_por_tipo_evento"][0]
        top_client = response.data["top_5_clientes"][0]
        self.assertEqual(metric["cantidad_proyectos"], 1)
        self.assertEqual(Decimal(str(metric["total_generado"])), Decimal("119.00"))
        self.assertEqual(
            Decimal(str(top_client["total_invertido"])), Decimal("119.00")
        )

    def test_report_exposes_pipeline_statuses_payments_and_monthly_activity(self):
        self.project.pagado = True
        self.project.fecha_pago = date.today()
        self.project.save(update_fields=["pagado", "fecha_pago"])
        Interaccion.objects.create(
            cliente=self.customer,
            tipo="llamada",
            resultado="contactado",
            descripcion="El cliente respondió la llamada",
            created_by=self.sales,
        )
        Cliente.objects.create(
            nombre="Cliente potencial",
            email="potential@example.com",
            estado="potencial",
        )
        self.client.force_authenticate(self.sales)

        response = self.client.get("/api/reports/gerencial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        pipeline = response.data["clientes_conversion"]
        self.assertEqual(pipeline["total_clientes"], 2)
        self.assertEqual(pipeline["clientes_que_cotizaron"], 1)
        self.assertEqual(pipeline["clientes_que_concretaron"], 1)
        self.assertEqual(pipeline["clientes_que_pagaron"], 1)
        self.assertEqual(
            {row["estado"]: row["total"] for row in response.data["cotizaciones_por_estado"]},
            {"aceptada": 1, "rechazada": 1},
        )
        self.assertEqual(
            response.data["interacciones_por_resultado"][0]["resultado"],
            "contactado",
        )
        self.assertEqual(
            response.data["clientes_por_proyecto"][0]["cliente_nombre"],
            "Cliente reportes",
        )
        current_month = date.today().strftime("%Y-%m")
        activity = next(
            row for row in response.data["actividad_mensual"] if row["periodo"] == current_month
        )
        self.assertEqual(activity["proyectos"], 1)
        self.assertEqual(activity["interacciones"], 1)
        self.assertEqual(activity["cotizaciones"], 2)

    def test_report_rejects_invalid_date_range(self):
        self.client.force_authenticate(self.sales)

        response = self.client.get(
            "/api/reports/gerencial/",
            {"desde": "2026-10-09", "hasta": "2026-10-08"},
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("periodo", response.data)

    def test_period_filter_keeps_existing_clients_in_conversion_denominator(self):
        Cliente.objects.filter(id=self.customer.id).update(
            created_at=date.today() - timedelta(days=90)
        )
        self.client.force_authenticate(self.sales)

        response = self.client.get(
            "/api/reports/gerencial/",
            {"desde": date.today().isoformat(), "hasta": date.today().isoformat()},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        pipeline = response.data["clientes_conversion"]
        self.assertEqual(pipeline["total_clientes"], 1)
        self.assertEqual(pipeline["clientes_que_cotizaron"], 1)
        self.assertLessEqual(pipeline["tasa_cliente_a_cotizacion"], 100)

    def test_paid_client_requires_an_accepted_quotation(self):
        customer = Cliente.objects.create(
            nombre="Pago sin cierre", email="paid-without-close@example.com"
        )
        Proyecto.objects.create(
            cliente=customer,
            nombre="Proyecto sin cotización aceptada",
            descripcion="No debe entrar al embudo pagado",
            tipo_evento="show",
            fecha_inicio=date.today(),
            fecha_fin=date.today() + timedelta(days=1),
            presupuesto_estimado=Decimal("300.00"),
            pagado=True,
            fecha_pago=date.today(),
            responsable=self.sales,
        )
        self.client.force_authenticate(self.sales)

        response = self.client.get("/api/reports/gerencial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        pipeline = response.data["clientes_conversion"]
        self.assertEqual(pipeline["clientes_que_concretaron"], 1)
        self.assertEqual(pipeline["clientes_que_pagaron"], 0)
        self.assertLessEqual(pipeline["tasa_cierre_a_pago"], 100)

    def test_paid_revenue_uses_latest_accepted_quote_once(self):
        self.project.pagado = True
        self.project.fecha_pago = date.today()
        self.project.save(update_fields=["pagado", "fecha_pago"])
        Cotizacion.objects.filter(projecto=self.project, estado="rechazada").update(
            estado="aceptada"
        )
        self.client.force_authenticate(self.sales)

        response = self.client.get("/api/reports/gerencial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            Decimal(str(response.data["resumen_financiero"]["ingresos_pagados"])),
            Decimal("1071.00"),
        )

    def test_payment_date_respects_period_filter(self):
        self.project.pagado = True
        self.project.fecha_pago = date.today() - timedelta(days=1)
        self.project.save(update_fields=["pagado", "fecha_pago"])
        self.client.force_authenticate(self.sales)

        response = self.client.get(
            "/api/reports/gerencial/",
            {"desde": date.today().isoformat(), "hasta": date.today().isoformat()},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["clientes_conversion"]["clientes_que_pagaron"], 0)

    def test_project_quotation_counts_respect_period_filter(self):
        Cotizacion.objects.filter(projecto=self.project).update(
            created_at=date.today() - timedelta(days=1)
        )
        self.client.force_authenticate(self.sales)

        response = self.client.get(
            "/api/reports/gerencial/",
            {"desde": date.today().isoformat(), "hasta": date.today().isoformat()},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["resumen_conversion"]["total"], 0)
        project_row = response.data["clientes_por_proyecto"][0]
        self.assertEqual(project_row["cotizaciones"], 0)
        self.assertEqual(project_row["cotizaciones_aceptadas"], 0)

    def test_responsible_filter_scopes_client_status_cohort(self):
        Cliente.objects.create(nombre="Fuera del equipo", email="outside@example.com")
        self.client.force_authenticate(self.sales)

        response = self.client.get(
            "/api/reports/gerencial/", {"responsable": self.sales.id}
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["clientes_conversion"]["total_clientes"], 1)
        self.assertEqual(sum(row["total"] for row in response.data["clientes_por_estado"]), 1)

    def test_project_filters_must_match_the_same_project_for_interactions(self):
        other_project = Proyecto.objects.create(
            cliente=self.customer,
            nombre="Proyecto con filtros separados",
            descripcion="No debe producir una coincidencia combinada",
            tipo_evento="concierto",
            fecha_inicio=date.today(),
            fecha_fin=date.today() + timedelta(days=1),
            presupuesto_estimado=Decimal("250.00"),
            estado="cancelado",
            responsable=self.sales,
        )
        self.project.estado = "aprobado"
        self.project.tipo_evento = "show"
        self.project.save(update_fields=["estado", "tipo_evento"])
        Interaccion.objects.create(
            cliente=self.customer,
            tipo="llamada",
            resultado="contactado",
            descripcion="Interacción del mismo cliente",
            created_by=self.sales,
        )
        self.client.force_authenticate(self.sales)

        response = self.client.get(
            "/api/reports/gerencial/",
            {"estado_proyecto": "aprobado", "tipo_evento": "concierto"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(other_project.estado, "cancelado")
        self.assertEqual(response.data["clientes_conversion"]["total_clientes"], 0)
        self.assertEqual(response.data["interacciones_por_resultado"], [])

    def test_production_cannot_bulk_export_managerial_report(self):
        production = User.objects.create_user(
            username="production-export-report",
            password="StrongPass!2026",
            role="production",
        )
        self.client.force_authenticate(production)

        response = self.client.get("/api/reports/gerencial/export/")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_managerial_report_export_contains_analysis_sheets(self):
        self.client.force_authenticate(self.sales)

        response = self.client.get("/api/reports/gerencial/export/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Resumen",
                "Proyectos y clientes",
                "Cotizaciones",
                "Interacciones",
                "Actividad mensual",
                "Actividad semanal",
                "Estados",
                "Resultados interacción",
            ],
        )
        self.assertEqual(workbook["Proyectos y clientes"]["B2"].value, "Evento")
