from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value):
    if isinstance(value, str) and value.startswith(DANGEROUS_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def export_quotations_xlsx(queryset):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cotizaciones"
    sheet.append(
        [
            "Número",
            "Proyecto",
            "Cliente",
            "Estado",
            "Fecha",
            "Vencimiento",
            "Subtotal",
            "Impuestos",
            "Total",
            "Notas",
        ]
    )
    fill = PatternFill("solid", fgColor="312E81")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill

    for quotation in queryset:
        sheet.append(
            [
                safe_excel_value(quotation.numero or ""),
                safe_excel_value(quotation.projecto.nombre),
                safe_excel_value(quotation.projecto.cliente.nombre),
                quotation.get_estado_display(),
                quotation.created_at,
                quotation.fecha_vencimiento,
                quotation.subtotal,
                quotation.impuestos,
                quotation.total,
                safe_excel_value(quotation.notas or ""),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {
        "A": 22,
        "B": 28,
        "C": 26,
        "D": 16,
        "E": 14,
        "F": 16,
        "G": 16,
        "H": 16,
        "I": 16,
        "J": 45,
    }.items():
        sheet.column_dimensions[column].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="historial_cotizaciones.xlsx"'
    workbook.save(response)
    return response
