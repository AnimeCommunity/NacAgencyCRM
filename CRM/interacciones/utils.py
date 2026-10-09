from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from django.http import HttpResponse


DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value):
    if isinstance(value, str) and value.startswith(DANGEROUS_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def export_interactions_xlsx(queryset):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Interacciones"
    headers = ["Cliente", "Tipo", "Resultado", "Descripción", "Fecha", "Responsable"]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="312E81")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    for interaction in queryset:
        sheet.append(
            [
                safe_excel_value(interaction.cliente.nombre),
                interaction.get_tipo_display(),
                interaction.get_resultado_display(),
                safe_excel_value(interaction.descripcion),
                interaction.created_at.replace(tzinfo=None),
                safe_excel_value(
                    interaction.created_by.get_full_name()
                    or interaction.created_by.username
                    if interaction.created_by
                    else "Sistema"
                ),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 24, "B": 24, "C": 22, "D": 55, "E": 22, "F": 24}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="historial_interacciones.xlsx"'
    workbook.save(response)
    return response
