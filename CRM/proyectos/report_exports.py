from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value):
    if isinstance(value, str) and value.startswith(DANGEROUS_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _style_sheet(sheet, widths):
    fill = PatternFill("solid", fgColor="312E81")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def export_management_report_xlsx(payload, querysets):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["Indicador", "Valor"])
    conversion = payload["resumen_conversion"]
    pipeline = payload["clientes_conversion"]
    financial = payload["resumen_financiero"]
    summary_rows = [
        ("Cotizaciones totales", conversion["total"]),
        ("Cotizaciones aceptadas", conversion["aceptadas"]),
        ("Cotizaciones pendientes", conversion["pendientes"]),
        ("Cotizaciones rechazadas", conversion["rechazadas"]),
        ("Cotizaciones vencidas", conversion["vencidas"]),
        ("Tasa de éxito (%)", conversion["tasa_exito_porcentaje"]),
        ("Clientes", pipeline["total_clientes"]),
        ("Clientes que cotizaron", pipeline["clientes_que_cotizaron"]),
        ("Clientes que concretaron", pipeline["clientes_que_concretaron"]),
        ("Clientes que pagaron", pipeline["clientes_que_pagaron"]),
        ("Ingresos aceptados", financial["ingresos_aceptados"]),
        ("Ingresos pagados", financial["ingresos_pagados"]),
        ("Saldo por cobrar", financial["saldo_por_cobrar"]),
    ]
    for row in summary_rows:
        summary.append(row)
    _style_sheet(summary, {"A": 34, "B": 20})

    projects_sheet = workbook.create_sheet("Proyectos y clientes")
    projects_sheet.append(
        [
            "Cliente",
            "Proyecto",
            "Estado",
            "Pagado",
            "Fecha inicio",
            "Fecha fin",
            "Responsable",
            "Presupuesto",
        ]
    )
    for project in querysets["projects"]:
        projects_sheet.append(
            [
                safe_excel_value(project.cliente.nombre),
                safe_excel_value(project.nombre),
                project.get_estado_display(),
                "Sí" if project.pagado else "No",
                project.fecha_inicio,
                project.fecha_fin,
                safe_excel_value(project.responsable.username if project.responsable else "Sin asignar"),
                project.presupuesto_estimado,
            ]
        )
    _style_sheet(
        projects_sheet,
        {"A": 26, "B": 30, "C": 18, "D": 12, "E": 16, "F": 16, "G": 22, "H": 18},
    )

    quotations_sheet = workbook.create_sheet("Cotizaciones")
    quotations_sheet.append(
        ["Número", "Proyecto", "Cliente", "Estado", "Fecha", "Vencimiento", "Subtotal", "Impuestos", "Total"]
    )
    for quotation in querysets["quotations"]:
        quotations_sheet.append(
            [
                safe_excel_value(quotation.numero or ""),
                safe_excel_value(quotation.projecto.nombre),
                safe_excel_value(quotation.projecto.cliente.nombre),
                quotation.get_estado_display(),
                quotation.created_at,
                quotation.fecha_vencimiento,
                quotation.subtotal,
                quotation.impuestos,
                quotation.total,
            ]
        )
    _style_sheet(
        quotations_sheet,
        {"A": 22, "B": 28, "C": 26, "D": 16, "E": 14, "F": 16, "G": 16, "H": 16, "I": 16},
    )

    interactions_sheet = workbook.create_sheet("Interacciones")
    interactions_sheet.append(["Cliente", "Tipo", "Resultado", "Descripción", "Fecha", "Responsable"])
    for interaction in querysets["interactions"]:
        interactions_sheet.append(
            [
                safe_excel_value(interaction.cliente.nombre),
                interaction.get_tipo_display(),
                interaction.get_resultado_display(),
                safe_excel_value(interaction.descripcion),
                interaction.created_at.replace(tzinfo=None),
                safe_excel_value(interaction.created_by.username if interaction.created_by else "Sistema"),
            ]
        )
    _style_sheet(interactions_sheet, {"A": 26, "B": 25, "C": 22, "D": 55, "E": 22, "F": 22})

    activity_sheet = workbook.create_sheet("Actividad mensual")
    activity_sheet.append(["Periodo", "Proyectos", "Cotizaciones", "Interacciones", "Ingresos aceptados"])
    for row in payload["actividad_mensual"]:
        activity_sheet.append(
            [
                row["periodo"],
                row["proyectos"],
                row["cotizaciones"],
                row["interacciones"],
                row["ingresos_aceptados"],
            ]
        )
    _style_sheet(activity_sheet, {"A": 16, "B": 14, "C": 16, "D": 16, "E": 22})

    weekly_sheet = workbook.create_sheet("Actividad semanal")
    weekly_sheet.append(["Semana", "Proyectos", "Cotizaciones", "Interacciones", "Ingresos aceptados"])
    for row in payload["actividad_semanal"]:
        weekly_sheet.append(
            [
                row["periodo"],
                row["proyectos"],
                row["cotizaciones"],
                row["interacciones"],
                row["ingresos_aceptados"],
            ]
        )
    _style_sheet(weekly_sheet, {"A": 16, "B": 14, "C": 16, "D": 16, "E": 22})

    states_sheet = workbook.create_sheet("Estados")
    states_sheet.append(["Entidad", "Estado", "Cantidad"])
    for entity, key in (
        ("Cliente", "clientes_por_estado"),
        ("Proyecto", "proyectos_por_estado"),
        ("Cotización", "cotizaciones_por_estado"),
    ):
        for row in payload[key]:
            states_sheet.append([entity, row["estado"], row["total"]])
    _style_sheet(states_sheet, {"A": 18, "B": 24, "C": 14})

    outcomes_sheet = workbook.create_sheet("Resultados interacción")
    outcomes_sheet.append(["Resultado", "Cantidad"])
    for row in payload["interacciones_por_resultado"]:
        outcomes_sheet.append([row["resultado"], row["total"]])
    _style_sheet(outcomes_sheet, {"A": 28, "B": 14})

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="informe_gerencial_crm.xlsx"'
    workbook.save(response)
    return response
